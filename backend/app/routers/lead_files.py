"""Bulk upload of leads from Excel / CSV (for example an Apollo export) and download of leads to Excel / CSV.

Every upload and download is recorded in the Activity Log (who, when, what, how many, and the reason for downloads),
and the Admin is notified of downloads.
Uploaded leads start at New. They get no welcome email, no notification each and no automatic task, so a file of
100,000 leads doesn't flood customers, the bell or the task list. Duplicates of open leads are skipped.
"""
import csv
import io
import re
from datetime import datetime, timezone
from decimal import Decimal, InvalidOperation

from fastapi import APIRouter, Depends, File, Form, HTTPException, Query, UploadFile
from fastapi.responses import Response
from sqlalchemy import func, insert, or_, select
from sqlalchemy.orm import Session, selectinload

from app import schemas
from app.constants import LEAD_OPEN_STAGES, LEAD_SOURCES, LEAD_STAGES
from app.database import get_db
from app.models import Account, ActivityLog, Lead, StageHistory, User
from app.security import acting_user, can_view, owner_for_new, require, visible
from app.serializers import lead_out
from app.services import log_activity, notify

router = APIRouter(prefix="/api/leads", tags=["Leads: upload and download"])

MAX_FILE_BYTES = 50 * 1024 * 1024
MAX_ROWS = 100_000
CHUNK = 1000

# Column headings we understand (compared without case, spaces or symbols). Apollo's export headings are included.
COLUMNS = {
    "salutation": ("salutation",),
    "first_name": ("firstname", "first"),
    "last_name": ("lastname", "last", "surname"),
    "full_name": ("name", "fullname", "contactname", "leadname"),
    "title": ("title", "jobtitle", "designation", "position"),
    "company": ("company", "companyname", "account", "accountname", "organization", "organisation"),
    "email": ("email", "emailaddress", "workemail", "businessemail"),
    "phone": ("phone", "workdirectphone", "mobilephone", "mobile", "corporatephone", "phonenumber", "directphone",
              "companyphone", "otherphone"),
    "industry": ("industry",),
    "linkedin_url": ("personlinkedinurl", "linkedinurl", "linkedin", "linkedinprofile", "leadlinkedinurl",
                     "profileurl"),
    "company_website": ("website", "companywebsite", "companyurl", "url", "domain"),
    "company_linkedin_url": ("companylinkedinurl", "companylinkedin"),
    "annual_revenue": ("annualrevenue", "revenue", "turnover"),
    "employees": ("employees", "noemployees", "employeecount", "noofemployees", "numberofemployees", "headcount", "companysize"),
    "location": ("location",),
    "city": ("city",),
    "state": ("state", "region"),
    "country": ("country",),
    "source": ("leadsource", "source"),
}
LOOKUP = {alias: field for field, aliases in COLUMNS.items() for alias in aliases}
LIMITS = {"first_name": 50, "last_name": 50, "company": 200, "title": 100, "industry": 100, "email": 255, "phone": 30}
EMAIL_RE = re.compile(r"^[^@\s,;]+@[^@\s,;]+\.[^@\s,;]+$")
SOURCE_ALIASES = {s.lower(): s for s in LEAD_SOURCES} | {"web": "Website", "linkedin": "Social Media",
                                                         "social": "Social Media", "apollo": "Other"}


def key(heading: str) -> str:
    return re.sub(r"[^a-z0-9]", "", (heading or "").lower().replace("#", "no"))


def read_rows(upload: UploadFile, content: bytes) -> tuple[list[str], list[list]]:
    name = (upload.filename or "").lower()
    if name.endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook
        try:
            book = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        except Exception:
            raise HTTPException(status_code=400, detail="This Excel file can't be opened. Save it again as .xlsx or .csv.")
        rows = book.worksheets[0].iter_rows(values_only=True)
        header = next(rows, None) or []
        data = []
        for row in rows:
            if row and any(v not in (None, "") for v in row):
                data.append(list(row))
                if len(data) > MAX_ROWS:
                    break
        return [str(h or "") for h in header], data
    if name.endswith((".csv", ".txt")):
        for encoding in ("utf-8-sig", "cp1252", "latin-1"):
            try:
                text = content.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        reader = csv.reader(io.StringIO(text))
        header = next(reader, None) or []
        data = []
        for row in reader:
            if any(v.strip() for v in row):
                data.append(row)
                if len(data) > MAX_ROWS:
                    break
        return header, data
    if name.endswith(".xls"):
        raise HTTPException(status_code=400, detail="Old .xls files aren't supported. Save it as .xlsx or .csv.")
    raise HTTPException(status_code=400, detail="Upload an Excel (.xlsx) or CSV (.csv) file.")


def text_of(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return " ".join(str(value).split())


def number(value: str) -> Decimal | None:
    """'$5,000,000' / '5M' / '1.2 B' / '250k' -> Decimal. Ranges like '51-200' take the upper number."""
    if not value:
        return None
    value = value.lower().replace(",", "").replace("$", "").replace("usd", "").strip()
    if "-" in value:
        value = value.split("-")[-1].strip()
    match = re.match(r"^(\d+(?:\.\d+)?)\s*([kmb])?\+?$", value)
    if not match:
        raise ValueError("not a number")
    amount = Decimal(match.group(1)) * {"k": 1_000, "m": 1_000_000, "b": 1_000_000_000}.get(match.group(2) or "", 1)
    return amount


def parse(header: list[str], rows: list[list]):
    """Turns the sheet into lead dicts. Returns (leads, errors, warnings, mapped headings, ignored headings).
    Only Last Name and Company are required: a wrong email, link or number doesn't skip the row (like HubSpot and
    Salesforce imports); that one value is left blank and listed as a warning so the salesperson can fill it in."""
    fields = [LOOKUP.get(key(h)) for h in header]
    mapped = sorted({h for h, f in zip(header, fields) if f})
    ignored = [h for h, f in zip(header, fields) if not f and h]
    if "last_name" not in fields and "full_name" not in fields:
        raise HTTPException(status_code=400, detail="The file needs a 'Last Name' (or 'Name') column.")
    if "company" not in fields:
        raise HTTPException(status_code=400, detail="The file needs a 'Company' column.")
    leads, errors, warnings = [], [], []
    for number_in_file, row in enumerate(rows, start=2):  # row 1 is the heading
        raw = {}
        for field, value in zip(fields, row):
            if field and field not in raw:
                raw[field] = text_of(value)
        if not raw.get("last_name") and raw.get("full_name"):
            parts = raw["full_name"].split(" ")
            raw["last_name"] = parts[-1]
            raw.setdefault("first_name", " ".join(parts[:-1]))
        if not raw.get("last_name"):
            errors.append({"row": number_in_file, "reason": "Last name is missing"})
            continue
        if not raw.get("company"):
            errors.append({"row": number_in_file, "reason": "Company is missing"})
            continue
        lead = {f: raw[f][:limit] or None for f, limit in LIMITS.items() if raw.get(f)}
        blank = []  # values left out of this row because they aren't valid
        email = (raw.get("email") or "").lower()
        if email and not EMAIL_RE.match(email):
            blank.append(f"email '{email}' is not valid")
            email = ""
        lead["email"] = email or None
        salutation = raw.get("salutation", "").title().rstrip(".") + "." if raw.get("salutation") else None
        lead["salutation"] = salutation if salutation in ("Mr.", "Ms.", "Mrs.", "Dr.", "Prof.") else None
        lead["source"] = SOURCE_ALIASES.get((raw.get("source") or "").lower(), "Upload")  # source from the file, or Upload
        location = raw.get("location") or ", ".join(x for x in (raw.get("city"), raw.get("state"), raw.get("country")) if x)
        lead["location"] = location[:200] or None
        for field, label, linkedin, limit in (("linkedin_url", "LinkedIn URL", True, 300),
                                              ("company_website", "company website", False, 255),
                                              ("company_linkedin_url", "company LinkedIn URL", True, 300)):
            value = raw.get(field) or None
            try:
                lead[field] = schemas.clean_url(value, linkedin=linkedin) if value else None
                if lead[field] and len(lead[field]) > limit:
                    raise ValueError("too long")
            except ValueError:
                lead[field] = None
                blank.append(f"{label} '{value[:60]}' is not valid")
        for field, label in (("annual_revenue", "turnover"), ("employees", "headcount")):
            try:
                value = number(raw.get(field, ""))
            except (ValueError, InvalidOperation):
                value = None
                blank.append(f"{label} '{raw.get(field)}' is not a number")
            if value is not None and value >= (Decimal("1e14") if field == "annual_revenue" else 10_000_000):
                value = None
            lead[field] = int(value) if field == "employees" and value is not None else value
        if blank:
            warnings.append({"row": number_in_file, "reason": "Uploaded; left blank: " + "; ".join(blank)})
        lead["_row"] = number_in_file
        leads.append(lead)
    return leads, errors, warnings, mapped, ignored


def drop_duplicates(db: Session, leads: list[dict], errors: list[dict]) -> list[dict]:
    """Skips rows whose email or LinkedIn profile is already on an open lead, or earlier in the same file."""
    emails = {l["email"] for l in leads if l.get("email")}
    profiles = {l["linkedin_url"].lower().rstrip("/") for l in leads if l.get("linkedin_url")}
    taken_emails, taken_profiles = set(), set()
    open_leads = Lead.stage.in_(LEAD_OPEN_STAGES)
    # one pass over the open leads (fast even with 100,000s of them), then compare in memory
    if emails:
        taken_emails = emails & set(db.scalars(select(func.lower(Lead.email)).where(
            open_leads, Lead.email.is_not(None))).all())
    if profiles:
        for url in db.scalars(select(Lead.linkedin_url).where(open_leads, Lead.linkedin_url.is_not(None))).all():
            if url.lower().rstrip("/") in profiles:
                taken_profiles.add(url.lower().rstrip("/"))
    keep, seen_emails, seen_profiles = [], set(), set()
    for lead in leads:
        email = lead.get("email")
        profile = (lead.get("linkedin_url") or "").lower().rstrip("/")
        if email and email in taken_emails or profile and profile in taken_profiles:
            errors.append({"row": lead["_row"], "reason": "Already an open lead (duplicate)", "duplicate": True})
        elif email and email in seen_emails or profile and profile in seen_profiles:
            errors.append({"row": lead["_row"], "reason": "Repeated in this file (duplicate)", "duplicate": True})
        else:
            keep.append(lead)
            seen_emails.add(email)
            seen_profiles.add(profile)
    return keep


@router.post("/import")
async def import_leads(
    file: UploadFile = File(...),
    owner_id: int | None = Form(None),
    dry_run: bool = Form(True),
    db: Session = Depends(get_db),
):
    """Upload leads from Excel / CSV. dry_run=true (default) only checks the file and shows what would happen."""
    me = require("import_leads")
    if me.role == "Admin":
        # the Admin doesn't own leads: the uploaded leads go to a Manager or Salesperson
        if owner_id is None or owner_id == me.id:
            raise HTTPException(status_code=400, detail="Choose who will own the uploaded leads.")
        target = db.get(User, owner_id)
    else:
        owner_id = owner_for_new(owner_id)
        target = db.get(User, owner_id)
    if target is None or not target.is_active or target.role == "Admin":
        raise HTTPException(status_code=400, detail="The owner must be an active Manager or Salesperson.")
    content = await file.read(MAX_FILE_BYTES + 1)
    if len(content) > MAX_FILE_BYTES:
        raise HTTPException(status_code=413, detail="The file is larger than 50 MB. Split it into smaller files.")
    header, rows = read_rows(file, content)
    if not rows:
        raise HTTPException(status_code=400, detail="The file has no rows to upload.")
    if len(rows) > MAX_ROWS:
        raise HTTPException(status_code=400, detail=f"Upload up to {MAX_ROWS:,} rows per file. Split the file and upload each part.")
    leads, errors, warnings, mapped, ignored = parse(header, rows)
    leads = drop_duplicates(db, leads, errors)
    kept = {l["_row"] for l in leads}
    warnings = [w for w in warnings if w["row"] in kept]
    errors.sort(key=lambda e: e["row"])
    skipped = sum(1 for e in errors if e.get("duplicate"))
    result = {
        "file": file.filename,
        "rows": len(rows),
        "ready": len(leads),
        "duplicates": skipped,
        "invalid": len(errors) - skipped,
        "owner": target.full_name,
        "columns_used": mapped,
        "columns_ignored": ignored,
        "errors": [{"row": e["row"], "reason": e["reason"]} for e in errors[:500]],
        "with_blanks": len(warnings),  # uploaded, but a wrong value was left blank
        "warnings": warnings[:500],
        "preview": [{k: (float(v) if isinstance(v, Decimal) else v) for k, v in l.items() if not k.startswith("_")}
                    for l in leads[:5]],
        "created": 0,
    }
    if dry_run or not leads:
        return result
    note = f"Uploaded from {file.filename}"
    for i in range(0, len(leads), CHUNK):
        batch = [{k: v for k, v in l.items() if not k.startswith("_")} | {"owner_id": owner_id, "stage": "New", "untouched": True,
                                                    "uploaded_from": (file.filename or "file")[:255]}
                 for l in leads[i:i + CHUNK]]
        ids = db.scalars(insert(Lead).returning(Lead.id), batch).all()
        db.execute(insert(StageHistory), [{"lead_id": lead_id, "from_stage": None, "to_stage": "New",
                                           "changed_by_id": me.id, "note": note} for lead_id in ids])
        result["created"] += len(ids)
    log_activity(db, me.id, "Leads uploaded", file.filename or "file",
                 f"{result['created']:,} created ({len(warnings):,} with fields left blank) • {skipped:,} duplicates skipped"
                 f" • {result['invalid']:,} invalid rows"
                 f" • Owner: {target.full_name}")
    if owner_id != me.id:
        notify(db, owner_id, "Leads assigned to you",
               f"{me.full_name} uploaded {result['created']:,} leads and assigned them to you ({file.filename}).",
               link="/leads", category="assignment")
    db.commit()
    return result


# ---------- the Leads list, page by page ----------

LIST_VIEWS = {
    "all_open": lambda q, me: q.where(Lead.stage.in_(LEAD_OPEN_STAGES)),
    "my_open": lambda q, me: q.where(Lead.stage.in_(LEAD_OPEN_STAGES), Lead.owner_id == me.id),
    "all": lambda q, me: q,
    "converted": lambda q, me: q.where(Lead.stage == "Converted"),
    "lost": lambda q, me: q.where(Lead.stage == "Lost"),
    "uploaded": lambda q, me: q.where(Lead.uploaded_from.is_not(None)),
    "my_uploaded": lambda q, me: q.where(Lead.uploaded_from.is_not(None), Lead.owner_id == me.id),
}
SORTS = {"name": Lead.name, "title": Lead.title, "company": Lead.company, "phone": Lead.phone, "email": Lead.email,
         "stage": Lead.stage, "source": Lead.source, "owner": User.full_name, "created": Lead.created_at}


def as_utc(value: datetime) -> datetime:
    """Dates from the browser come with their time zone; lead times are stored in UTC without one."""
    return value.astimezone(timezone.utc).replace(tzinfo=None) if value.tzinfo else value


@router.get("/page")
def leads_page(
    view: str = Query("all_open", description=", ".join(LIST_VIEWS)),
    stage: str | None = Query(None, description="Only this Lead Status (e.g. Demo)"),
    source: str | None = Query(None, description="Only this Lead Source (e.g. Upload)"),
    owner: str | None = Query(None, description="Owner user id, or 'me'"),
    created_from: datetime | None = Query(None, description="Created on or after (ISO date/time)"),
    created_to: datetime | None = Query(None, description="Created before (ISO date/time)"),
    search: str | None = Query(None, description="Name, company, title, email, phone or location"),
    sort: str = Query("created", description=", ".join(SORTS)),
    direction: str = Query("desc", pattern="^(asc|desc)$"),
    offset: int = Query(0, ge=0),
    limit: int = Query(100, ge=1, le=500),
    db: Session = Depends(get_db),
):
    """The Leads list views, loaded page by page from the database (there can be 100,000s of leads),
    including uploaded ones. Returns the total and one page."""
    me = acting_user()
    if view not in LIST_VIEWS:
        raise HTTPException(status_code=400, detail="Unknown list view")
    if sort not in SORTS:
        raise HTTPException(status_code=400, detail="Unknown sort column")
    query = LIST_VIEWS[view](visible(select(Lead), Lead), me)
    if stage:
        query = query.where(Lead.stage == stage)
    if source:
        query = query.where(Lead.source == source)
    if owner:
        if owner == "me":
            query = query.where(Lead.owner_id == me.id)
        elif owner.isdigit():
            query = query.where(Lead.owner_id == int(owner))
        else:
            raise HTTPException(status_code=400, detail="Unknown owner")
    if created_from:
        query = query.where(Lead.created_at >= as_utc(created_from))
    if created_to:
        query = query.where(Lead.created_at < as_utc(created_to))
    if search and search.strip():
        like = f"%{search.strip()}%"
        query = query.where(or_(Lead.name.ilike(like), Lead.company.ilike(like), Lead.email.ilike(like),
                                Lead.title.ilike(like), Lead.phone.ilike(like), Lead.location.ilike(like)))
    total = db.scalar(select(func.count()).select_from(query.subquery()))
    column = SORTS[sort]
    if sort == "owner":
        query = query.join(User, User.id == Lead.owner_id)
    order = column.desc().nulls_last() if direction == "desc" else column.asc().nulls_last()
    rows = db.scalars(query.options(selectinload(Lead.owner), selectinload(Lead.opportunity))
                      .order_by(order, Lead.id.desc()).offset(offset).limit(limit)).all()
    return {"total": total, "offset": offset, "limit": limit, "items": [lead_out(l) for l in rows]}


# ---------- same company already in the CRM (shown while typing the company on a lead) ----------

COMPANY_SUFFIXES = {"inc", "incorporated", "llc", "ltd", "limited", "pvt", "private", "corp", "corporation", "co",
                    "company", "gmbh", "plc", "llp", "lp", "sa", "ag", "bv", "pte", "the"}


def company_key(name: str | None) -> str:
    """'Brightline Logistics, Inc.' -> 'brightline logistics' (case, punctuation and company suffixes ignored)."""
    words = re.sub(r"[^a-z0-9& ]", " ", (name or "").lower()).split()
    return " ".join(w for w in words if w not in COMPANY_SUFFIXES)


@router.get("/company-matches")
def company_matches(
    company: str = Query(..., min_length=2, max_length=200),
    exclude_id: int | None = Query(None, description="The lead being edited"),
    db: Session = Depends(get_db),
):
    """Leads and accounts that already have this company. Records the user can open come with their details and
    a link; others show only the owner and stage, so the salesperson knows whose team to check with."""
    wanted = company_key(company)
    if len(wanted) < 2:
        return {"company": company, "total": 0, "matches": []}
    first_word = wanted.split()[0]
    like = f"%{first_word}%"
    matches = []
    leads = db.scalars(select(Lead).options(selectinload(Lead.owner)).where(Lead.company.ilike(like))
                       .order_by(Lead.updated_at.desc()).limit(5000)).all()
    for lead in leads:
        if lead.id == exclude_id or company_key(lead.company) != wanted:
            continue
        open_lead = lead.stage in LEAD_OPEN_STAGES
        visible_to_me = can_view(lead.owner_id)
        matches.append({
            "type": "Lead", "id": lead.id if visible_to_me else None, "name": lead.name if visible_to_me else None,
            "company": lead.company, "stage": lead.stage, "open": open_lead, "owner": lead.owner.full_name if lead.owner else None,
            "owner_id": lead.owner_id, "mine": lead.owner_id == acting_user().id, "can_open": visible_to_me,
        })
    accounts = db.scalars(select(Account).options(selectinload(Account.owner)).where(Account.name.ilike(like)).limit(200)).all()
    for account in accounts:
        if company_key(account.name) != wanted:
            continue
        visible_to_me = can_view(account.owner_id)
        matches.append({
            "type": "Account", "id": account.id if visible_to_me else None, "name": account.name,
            "company": account.name, "stage": "Customer account", "open": True,
            "owner": account.owner.full_name if account.owner else None, "owner_id": account.owner_id,
            "mine": account.owner_id == acting_user().id, "can_open": visible_to_me,
        })
    matches.sort(key=lambda m: (m["type"] != "Account", not m["open"], not m["mine"]))
    owners: dict[str, int] = {}
    for m in matches:
        name = "You" if m["mine"] else (m["owner"] or "Unknown")
        owners[name] = owners.get(name, 0) + 1
    return {"company": company, "total": len(matches), "matches": matches[:10],
            "owners": [{"owner": o, "count": c} for o, c in sorted(owners.items(), key=lambda x: -x[1])]}


# ---------- upload and download history ----------

TRANSFER_ACTIONS = {
    "leads": ("Leads uploaded", "Leads downloaded"),
    "opportunities": ("Opportunities uploaded", "Opportunities downloaded"),
}


@router.get("/transfers")
def transfer_history(
    object: str = Query("leads", pattern="^(leads|opportunities)$"),
    limit: int = Query(200, ge=1, le=1000),
    db: Session = Depends(get_db),
):
    """Who uploaded or downloaded leads / opportunities, when, what and why (newest first).
    The Admin and Managers see everyone's; other users see their own."""
    me = acting_user()
    query = select(ActivityLog).where(ActivityLog.action.in_(TRANSFER_ACTIONS[object]))
    if not (me.role in ("Admin", "Manager") or me.can("manage_users")):
        query = query.where(ActivityLog.actor_id == me.id)
    rows = db.scalars(query.options(selectinload(ActivityLog.actor))
                      .order_by(ActivityLog.created_at.desc(), ActivityLog.id.desc()).limit(limit)).all()
    return [{"id": r.id, "at": r.created_at, "by": r.actor.full_name if r.actor else "System",
             "action": "Upload" if r.action.endswith("uploaded") else "Download", "what": r.target,
             "details": r.details} for r in rows]


# ---------- download ----------

EXPORT_COLUMNS = (
    ("Salutation", "salutation"), ("First Name", "first_name"), ("Last Name", "last_name"), ("Title", "title"),
    ("Company", "company"), ("Email", "email"), ("Phone", "phone"), ("LinkedIn URL", "linkedin_url"),
    ("Company Website", "company_website"), ("Company LinkedIn URL", "company_linkedin_url"),
    ("Turnover (USD)", "annual_revenue"), ("Headcount", "employees"), ("Location", "location"),
    ("Industry", "industry"), ("Lead Source", "source"), ("Status", "stage"), ("Owner", "owner"),
    ("Lost Reason", "lost_reason"), ("Created", "created_at"), ("Converted", "converted_at"),
)
STAGE_FILTERS = {"all": None, "open": LEAD_OPEN_STAGES, **{s.lower(): (s,) for s in LEAD_STAGES}}


@router.get("/export")
def export_leads(
    reason: str = Query(..., min_length=5, max_length=300, description="Why the file is needed (recorded)"),
    format: str = Query("xlsx", pattern="^(xlsx|csv)$"),
    stage: str = Query("all", description="all, open, new, nurturing, demo, converted or lost"),
    mine: bool = Query(False, description="Only my leads"),
    db: Session = Depends(get_db),
):
    """Download the leads you can see to Excel / CSV. The reason, filters and row count go in the Activity Log
    and the Admin is notified."""
    me = require("export_data")
    reason = " ".join(reason.split())
    if len(reason) < 5:
        raise HTTPException(status_code=422, detail="Give a reason for the download (at least 5 characters).")
    stage = stage.lower()
    if stage not in STAGE_FILTERS:
        raise HTTPException(status_code=400, detail="Unknown status filter")
    query = visible(select(Lead), Lead).order_by(Lead.created_at.desc(), Lead.id.desc())
    if STAGE_FILTERS[stage]:
        query = query.where(Lead.stage.in_(STAGE_FILTERS[stage]))
    if mine:
        query = query.where(Lead.owner_id == me.id)
    owners = dict(db.execute(select(User.id, User.full_name)).all())
    leads = db.scalars(query).all()

    def values(lead: Lead) -> list:
        out = []
        for _, attr in EXPORT_COLUMNS:
            value = owners.get(lead.owner_id) if attr == "owner" else getattr(lead, attr)
            if isinstance(value, Decimal):
                value = float(value)
            if isinstance(value, datetime):
                value = value.strftime("%Y-%m-%d %H:%M")
            out.append(value)
        return out

    stamp = datetime.now().strftime("%Y%m%d-%H%M")
    label = {"all": "All", "open": "Open"}.get(stage, stage.title())
    filename = f"SalesTrack-Leads-{label}{'-Mine' if mine else ''}-{stamp}.{format}"
    headings = [h for h, _ in EXPORT_COLUMNS]
    if format == "csv":
        buffer = io.StringIO()
        writer = csv.writer(buffer)
        writer.writerow(headings)
        writer.writerows(values(l) for l in leads)
        body, media = ("﻿" + buffer.getvalue()).encode(), "text/csv; charset=utf-8"
    else:
        from openpyxl import Workbook
        from openpyxl.styles import Font
        book = Workbook(write_only=True)
        sheet = book.create_sheet("Leads")
        from openpyxl.cell import WriteOnlyCell
        bold = []
        for h in headings:
            cell = WriteOnlyCell(sheet, value=h)
            cell.font = Font(bold=True)
            bold.append(cell)
        sheet.append(bold)
        for lead in leads:
            sheet.append(values(lead))
        out = io.BytesIO()
        book.save(out)
        body, media = out.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    what = f"{len(leads):,} {'lead' if len(leads) == 1 else 'leads'} ({label}{', mine only' if mine else ''})"
    log_activity(db, me.id, "Leads downloaded", what, f"Reason: {reason} • {filename}")
    admins = db.scalars(select(User.id).where(User.role == "Admin", User.is_active.is_(True), User.id != me.id)).all()
    for admin_id in admins:
        notify(db, admin_id, "Leads downloaded", f"{me.full_name} downloaded {what}.\nReason: {reason}",
               link="/activity")
    db.commit()
    return Response(content=body, media_type=media,
                    headers={"Content-Disposition": f'attachment; filename="{filename}"',
                             "Access-Control-Expose-Headers": "Content-Disposition"})
