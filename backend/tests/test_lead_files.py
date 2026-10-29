"""Lead links and company size, deleting leads, and bulk upload / download of leads (with tracking)."""
import csv
import io

import pytest
from openpyxl import Workbook, load_workbook

from app import mailer
from tests.conftest import login_as


@pytest.fixture
def outbox(monkeypatch):
    sent = []
    monkeypatch.setattr(mailer, "send", lambda to, subject, body, **kw: sent.append({"to": to, "subject": subject}))
    return sent


def new_lead(client, **extra):
    body = {"last_name": "Menon", "first_name": "Priya", "company": "Brightline", "welcome_email": False} | extra
    res = client.post("/api/leads", json=body)
    assert res.status_code == 201, res.text
    return res.json()


def csv_file(rows, name="apollo.csv"):
    out = io.StringIO()
    csv.writer(out).writerows(rows)
    return {"file": (name, out.getvalue().encode(), "text/csv")}


APOLLO = [
    ["First Name", "Last Name", "Title", "Company", "Email", "Person Linkedin Url", "Website",
     "Company Linkedin Url", "# Employees", "Annual Revenue", "City", "State", "Country", "Seniority"],
    ["Asha", "Iyer", "CTO", "Nimbus", "asha@nimbus.example", "linkedin.com/in/asha", "nimbus.example",
     "https://www.linkedin.com/company/nimbus", "250", "$12,500,000", "Austin", "Texas", "United States", "C-Suite"],
    ["Ben", "Cole", "VP Sales", "Orbit", "ben@orbit.example", "", "", "", "51-200", "3M", "", "", "India", ""],
    ["", "", "Nobody", "Nameless", "", "", "", "", "", "", "", "", "", ""],          # no last name
    ["Cara", "Diaz", "CEO", "Acme", "not-an-email", "", "", "", "", "", "", "", "", ""],  # bad email
    ["Asha", "Iyer", "CTO", "Nimbus", "ASHA@nimbus.example", "", "", "", "", "", "", "", "", ""],  # repeated
]


def test_lead_links_are_saved_cleaned_and_copied_on_convert(client, outbox):
    lead = new_lead(client, linkedin_url="linkedin.com/in/priya", company_website="brightline.example",
                    company_linkedin_url="www.linkedin.com/company/brightline", annual_revenue=5000000,
                    employees=120, location="Hyderabad, India")
    assert lead["linkedin_url"] == "https://linkedin.com/in/priya"
    assert lead["company_website"] == "https://brightline.example"
    assert lead["annual_revenue"] == 5000000 and lead["employees"] == 120
    bad = client.post("/api/leads", json={"last_name": "X", "company": "Y", "linkedin_url": "twitter.com/x"})
    assert bad.status_code == 422
    res = client.post(f"/api/leads/{lead['id']}/convert", json={"completed_stages": ["Nurturing", "Demo"]}).json()
    assert res["account"]["website"] == "https://brightline.example" and res["account"]["employees"] == 120
    assert res["contact"]["linkedin_url"] == "https://linkedin.com/in/priya"


def test_delete_lead_by_owner_manager_admin_only(client, outbox):
    ravis = new_lead(client, email="del@x.example")
    other = new_lead(login_as(client, 1), last_name="Mgr", email="m@x.example")  # Priya (Manager) owns this
    assert login_as(client, 2).delete(f"/api/leads/{other['id']}").status_code in (403, 404)
    client.post("/api/tasks", json={"title": "Call", "due_date": "2030-01-01T10:00", "lead_id": ravis["id"]})
    assert client.delete(f"/api/leads/{ravis['id']}").status_code == 200
    assert client.get(f"/api/leads/{ravis['id']}").status_code == 404
    bell = [n["title"] for n in login_as(client, 3).get("/api/notifications").json()]
    assert "Lead deleted" in bell
    log = client.get("/api/activity-log").json()[0]
    assert log["action"] == "Lead deleted" and log["target"] == "Priya Menon" and "Owner: Ravi Kumar" in log["details"]
    assert client.delete(f"/api/leads/{other['id']}").status_code == 200  # the Admin can delete anyone's
    # converted leads stay
    conv = new_lead(login_as(client, 2), email="c@x.example")
    client.post(f"/api/leads/{conv['id']}/convert", json={"completed_stages": ["Nurturing", "Demo"]})
    assert client.delete(f"/api/leads/{conv['id']}").status_code == 400


def test_upload_preview_then_import(client, outbox):
    new_lead(client, email="ben@orbit.example")  # Ben is already an open lead
    preview = client.post("/api/leads/import", files=csv_file(APOLLO), data={"dry_run": "true"}).json()
    # Cara's email is wrong: she is still uploaded, with the email left blank for the salesperson to fill in
    assert preview["rows"] == 5 and preview["ready"] == 2 and preview["duplicates"] == 2 and preview["invalid"] == 1
    assert preview["with_blanks"] == 1 and "email 'not-an-email' is not valid" in preview["warnings"][0]["reason"]
    assert "Seniority" in preview["columns_ignored"]
    assert preview["preview"][0]["location"] == "Austin, Texas, United States"
    assert preview["preview"][0]["employees"] == 250 and preview["preview"][0]["annual_revenue"] == 12500000
    assert client.get("/api/leads").json().__len__() == 1  # nothing saved yet
    sent = len(outbox)
    done = client.post("/api/leads/import", files=csv_file(APOLLO), data={"dry_run": "false"}).json()
    assert done["created"] == 2
    assert [l["last_name"] for l in client.get("/api/leads").json()] == ["Menon"]  # uploads stay out of the lists
    page = client.get("/api/leads/page", params={"view": "uploaded"}).json()
    assert page["total"] == 2
    asha = next(l for l in page["items"] if l["last_name"] == "Iyer")
    assert asha["uploaded_from"] == "apollo.csv"
    # uploaded leads are in the normal open-leads list too
    assert {l["last_name"] for l in client.get("/api/leads/page").json()["items"]} == {"Menon", "Iyer", "Diaz"}
    assert asha["linkedin_url"] == "https://linkedin.com/in/asha" and asha["owner_id"] == 2 and asha["stage"] == "New"
    assert len(outbox) == sent  # no welcome emails for uploaded leads
    assert client.get(f"/api/leads/{asha['id']}").json()["tasks"] == []  # and no automatic tasks
    log = login_as(client, 3).get("/api/activity-log").json()[0]
    assert log["action"] == "Leads uploaded" and log["by"] == "Ravi Kumar" and "2 created (1 with fields left blank)" in log["details"]


def test_upload_excel_and_admin_must_pick_owner(client, outbox):
    book = Workbook()
    book.active.append(["Name", "Company", "Email"])
    book.active.append(["Dev Patel", "Zen Labs", "dev@zen.example"])
    data = io.BytesIO()
    book.save(data)
    files = {"file": ("leads.xlsx", data.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    login_as(client, 3)
    assert client.post("/api/leads/import", files=files, data={"dry_run": "false"}).status_code == 400
    done = client.post("/api/leads/import", files=files, data={"dry_run": "false", "owner_id": "2"}).json()
    assert done["created"] == 1
    lead = login_as(client, 2).get("/api/leads/page", params={"view": "my_uploaded"}).json()["items"][0]
    assert (lead["first_name"], lead["last_name"], lead["owner_id"]) == ("Dev", "Patel", 2)
    assert "Leads assigned to you" in [n["title"] for n in client.get("/api/notifications").json()]
    # a salesperson can't upload for someone else
    assert client.post("/api/leads/import", files=files, data={"owner_id": "1"}).status_code == 403
    assert client.post("/api/leads/import", files={"file": ("a.pdf", b"x", "application/pdf")}).status_code == 400


def test_download_needs_reason_is_tracked_and_respects_access(client, outbox):
    new_lead(client, email="mine@x.example", employees=10)
    new_lead(login_as(client, 1), last_name="Other", email="other@x.example")  # Priya's lead
    login_as(client, 2)
    assert client.get("/api/leads/export").status_code == 422  # reason required
    res = client.get("/api/leads/export", params={"reason": "Weekly follow-up list", "format": "xlsx"})
    assert res.status_code == 200 and "attachment" in res.headers["content-disposition"]
    sheet = load_workbook(io.BytesIO(res.content)).active
    rows = list(sheet.iter_rows(values_only=True))
    assert rows[0][2] == "Last Name" and len(rows) == 2 and rows[1][2] == "Menon"  # only Ravi's own lead
    csv_res = client.get("/api/leads/export", params={"reason": "Backup copy", "format": "csv", "stage": "open"})
    assert csv_res.text.lstrip("﻿").startswith("Salutation,First Name")
    login_as(client, 3)
    log = client.get("/api/activity-log").json()
    assert log[0]["action"] == "Leads downloaded" and "Reason: Backup copy" in log[0]["details"]
    assert log[1]["target"].startswith("1 lead (") and log[1]["by"] == "Ravi Kumar"
    assert "Leads downloaded" in [n["title"] for n in client.get("/api/notifications").json()]
    # the Admin can switch downloads off for salespeople
    client.put("/api/permissions", json={"role": "Salesperson", "permission": "export_data", "allowed": False})
    assert login_as(client, 2).get("/api/leads/export", params={"reason": "Again please"}).status_code == 403


def test_leads_list_page_views_search_and_sort(client, outbox):
    rows = [["Last Name", "Company", "Title"]] + [[f"Lead{i}", f"Co{i}", "CTO" if i % 2 else "CFO"] for i in range(5)]
    client.post("/api/leads/import", files=csv_file(rows), data={"dry_run": "false"})
    new_lead(client, last_name="Manual", email="man@x.example")
    page = client.get("/api/leads/page", params={"limit": 2}).json()
    assert page["total"] == 6 and len(page["items"]) == 2
    assert client.get("/api/leads/page", params={"view": "uploaded"}).json()["total"] == 5
    assert client.get("/api/leads/page", params={"search": "cfo"}).json()["total"] == 3
    names = [l["last_name"] for l in client.get("/api/leads/page", params={"sort": "name", "direction": "asc"}).json()["items"]]
    assert names == sorted(names)
    first = client.get("/api/leads/page", params={"view": "uploaded"}).json()["items"][0]["id"]
    client.put(f"/api/leads/{first}/stage", json={"stage": "Nurturing"})
    assert client.get("/api/leads/page", params={"stage": "Nurturing"}).json()["total"] == 1
    # filters: Lead Source, Owner and Created date
    assert client.get("/api/leads/page", params={"source": "Upload"}).json()["total"] == 5
    assert client.get("/api/leads/page", params={"owner": "me"}).json()["total"] == 6
    assert client.get("/api/leads/page", params={"owner": "1"}).json()["total"] == 0
    assert client.get("/api/leads/page", params={"created_from": "2000-01-01T00:00:00Z"}).json()["total"] == 6
    assert client.get("/api/leads/page", params={"created_to": "2000-01-01T00:00:00Z"}).json()["total"] == 0
    assert client.get("/api/leads/page", params={"owner": "x"}).status_code == 400
    # a Manager sees everyone's leads; another salesperson's are hidden from Ravi
    assert login_as(client, 1).get("/api/leads/page", params={"view": "my_open"}).json()["total"] == 0
    assert login_as(client, 1).get("/api/leads/page").json()["total"] == 6



DEALS = [
    ["Opportunity Name", "Company", "Amount", "Close Date", "Stage", "Contact Name", "Contact Email", "Lost Reason"],
    ["Nimbus ERP rollout", "Nimbus", "$45,000", "2026-12-15", "Proposal", "Asha Iyer", "asha@nimbus.example", ""],
    ["Orbit renewal", "Orbit", "12000", "11/30/2026", "Won", "", "", ""],
    ["Zen pilot", "Zen Labs", "5k", "", "Closed Lost", "", "", "No budget"],
    ["", "", "100", "", "", "", "", ""],                                    # nothing to name it
    ["Bad stage", "Acme", "1", "", "Negotiation", "", "", ""],              # unknown stage
    ["Nimbus ERP rollout", "Nimbus", "1", "", "", "", "", ""],              # repeated
]


def test_opportunity_upload_and_download(client, outbox):
    preview = client.post("/api/opportunities/import", files=csv_file(DEALS, "deals.csv"), data={"dry_run": "true"}).json()
    assert (preview["ready"], preview["duplicates"], preview["invalid"]) == (3, 1, 2)
    assert preview["preview"][0]["contact"] == "Asha Iyer"
    done = client.post("/api/opportunities/import", files=csv_file(DEALS, "deals.csv"), data={"dry_run": "false"}).json()
    assert done["created"] == 3 and done["accounts_created"] == 3
    opps = {o["name"]: o for o in client.get("/api/opportunities").json()}
    assert opps["Nimbus ERP rollout"]["stage"] == "Proposal" and opps["Nimbus ERP rollout"]["probability"] == 75
    assert opps["Orbit renewal"]["final_value"] == 12000 and opps["Zen pilot"]["lost_reason"] == "No budget"
    again = client.post("/api/opportunities/import", files=csv_file(DEALS, "deals.csv"), data={"dry_run": "true"}).json()
    assert again["ready"] == 0 and again["duplicates"] == 4  # everything is in the CRM now
    res = client.get("/api/opportunities/export", params={"reason": "Won deals for the review", "stage": "won"})
    rows = list(load_workbook(io.BytesIO(res.content)).active.iter_rows(values_only=True))
    assert rows[0][0] == "Opportunity Name" and [r[0] for r in rows[1:]] == ["Orbit renewal"]
    assert client.get("/api/opportunities/export").status_code == 422  # a reason is required


def test_transfer_history_on_the_list_pages(client, outbox):
    client.post("/api/leads/import", files=csv_file([["Last Name", "Company"], ["Roy", "Dune"]]), data={"dry_run": "false"})
    client.get("/api/leads/export", params={"reason": "Call list for today"})
    mine = client.get("/api/leads/transfers").json()
    assert [(h["action"], h["by"]) for h in mine] == [("Download", "Ravi Kumar"), ("Upload", "Ravi Kumar")]
    assert "Reason: Call list for today" in mine[0]["details"]
    assert client.get("/api/leads/transfers", params={"object": "opportunities"}).json() == []
    # the Manager sees everyone's; another salesperson only their own
    assert len(login_as(client, 1).get("/api/leads/transfers").json()) == 2


def test_company_matches_while_typing(client, outbox):
    mine = new_lead(client, company="Brightline Logistics", email="a@b.example")
    new_lead(login_as(client, 1), last_name="Mgr", company="Brightline Logistics, Inc.", email="m@b.example")  # Priya's
    new_lead(client, last_name="Other", company="Brightline Foods", email="o@b.example")  # a different company
    res = login_as(client, 2).get("/api/leads/company-matches", params={"company": "brightline logistics inc"}).json()
    assert res["total"] == 2 and {o["owner"] for o in res["owners"]} == {"You", "Priya Sharma"}
    own = next(m for m in res["matches"] if m["mine"])
    other = next(m for m in res["matches"] if not m["mine"])
    assert own["name"] == "Priya Menon" and own["can_open"] and own["id"] == mine["id"]
    assert other["owner"] == "Priya Sharma" and other["name"] is None and not other["can_open"]  # no details
    assert login_as(client, 1).get("/api/leads/company-matches", params={"company": "Brightline Logistics"}).json()["matches"][0]["can_open"]
    assert client.get("/api/leads/company-matches", params={"company": "Brightline Logistics", "exclude_id": mine["id"]}).json()["total"] == 1
